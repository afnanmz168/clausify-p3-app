"""
App features that sit around the models:

  1. extract_text()        — read .txt, .pdf and .docx uploads (locally, nothing is sent anywhere)
  2. highlighted_html()    — the whole contract with every located clause marked in its risk colour
  3. results_csv() / results_pdf() — downloadable risk report
  4. (the missing-protection checklist logic lives in model_utils.missing_protections)
"""
import csv
import io
from datetime import datetime
from html import escape as _html_escape


def escape(s):
    """HTML-escape text that will go through st.markdown: also neutralize '$' (Streamlit
    renders $...$ as maths) and turn newlines into <br> (a blank line would end the HTML
    block and let Markdown re-format the rest, e.g. "11. ..." becoming a list)."""
    return _html_escape(str(s)).replace("$", "&#36;").replace("\n", "<br>")

RISK_COLOURS = {"High": "#d94f5c", "Medium": "#e8a13a", "Low": "#2f9e6e", "Unrecognized": "#8a93b8"}
RISK_RANK = {"High": 0, "Medium": 1, "Low": 2, "Unrecognized": 3}


# --------------------------------------------------------------------------- #
# 1. File upload: .txt / .pdf / .docx
# --------------------------------------------------------------------------- #
def extract_text(name, data):
    """Return (text, note). note is a warning string or ""."""
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "txt":
        for enc in ("utf-8-sig", "cp1252"):          # UTF-8 (with/without BOM), then Windows text
            try:
                return data.decode(enc), ""
            except UnicodeDecodeError:
                pass
        return data.decode("latin-1"), ""             # never fails
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
        text = "\n\n".join(p.strip() for p in pages if p.strip())
        if len(text) < 50:
            return text, ("Very little text could be read from this PDF. It may be a scanned image; "
                          "scanned PDFs need OCR, which this app does not do.")
        return text, ""
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:                      # clauses are sometimes laid out in tables
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(" | ".join(cells))
        return "\n\n".join(paras), ""
    raise ValueError(f"Unsupported file type: .{ext}")


# --------------------------------------------------------------------------- #
# 2. Highlighted contract view
# --------------------------------------------------------------------------- #
def _spans(text, items):
    """Non-overlapping (start, end, item, anchor) spans; on overlap the higher risk wins."""
    cands = []
    for k, it in enumerate(items):
        quote = it.get("quote", "")
        start = it.get("start", -1)
        if quote and start is not None and start >= 0:
            cands.append((start, start + len(quote), it, f"clause-{k}"))
    cands.sort(key=lambda c: (RISK_RANK[c[2]["risk"]], c[0]))
    taken, out = [], []
    for s, e, it, a in cands:
        if all(e <= ts or s >= te for ts, te in taken):
            taken.append((s, e)); out.append((s, e, it, a))
    return sorted(out, key=lambda c: c[0])


def anchors(text, items):
    """{id(item): anchor} for items that are highlighted in the contract view."""
    return {id(it): a for _, _, it, a in _spans(text, items)}


def highlighted_html(text, items, max_chars=60000):
    """HTML of the contract with located clauses marked; long contracts are truncated."""
    shown = text[:max_chars]
    parts, pos = [], 0
    for s, e, it, a in _spans(shown, items):
        if s < pos or e > len(shown):
            continue
        cats = it.get("label", it.get("category") or "Unrecognized")
        colour = RISK_COLOURS[it["risk"]]
        parts.append(escape(shown[pos:s]))
        parts.append(
            f'<mark id="{a}" class="hl" title="{escape(cats)} — {it["risk"]} risk" '
            f'style="background:{colour}33;border-bottom:2px solid {colour};color:inherit">'
            f'{escape(shown[s:e])}<sup class="hl-tag" style="color:{colour}"> {escape(cats)}</sup></mark>')
        pos = e
    parts.append(escape(shown[pos:]))
    more = (f'\n\n… ({len(text) - max_chars:,} more characters not shown)' if len(text) > max_chars else "")
    return "".join(parts) + escape(more)


# --------------------------------------------------------------------------- #
# 3. Downloadable report
# --------------------------------------------------------------------------- #
CSV_FIELDS = ["#", "category", "risk", "status", "risk_reason", "plain_english", "summary_mismatch", "decided_by",
              "score", "transformer_score", "tfidf_score", "quoted_text"]


def results_csv(rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=CSV_FIELDS, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in CSV_FIELDS})
    return buf.getvalue().encode("utf-8-sig")       # BOM so Excel reads UTF-8 correctly


def results_pdf(rows, missing, meta):
    """Render the report as a PDF with PyMuPDF's HTML engine."""
    import fitz
    counts = {lvl: sum(1 for r in rows if r["risk"] == lvl) for lvl in RISK_COLOURS}
    cards = []
    for r in rows:
        c = RISK_COLOURS[r["risk"]]
        extra = f'<p class="pe"><b>In plain English:</b> {escape(r["plain_english"])}</p>' if r.get("plain_english") else ""
        if r.get("summary_mismatch"):
            extra += (f'<p class="sc">Note: the summarizer read this text as {escape(r["summary_mismatch"])}, '
                      f'not {escape(r["category"])}.</p>')
        scores = f'<p class="sc">{escape(r.get("score_text", ""))}</p>' if r.get("score_text") else ""
        cards.append(
            f'<div class="card" style="border-left:4px solid {c}">'
            f'<p class="h"><b>{escape(str(r["#"]))}. {escape(r["category"])}</b> '
            f'<span style="color:{c}"><b>{r["risk"].upper()}</b></span>'
            f'{" · <b>POSSIBLE — CHECK</b>" if r.get("possible") else ""}</p>'
            f'<p class="why">{escape(r.get("risk_reason", ""))}</p>{extra}{scores}'
            f'<p class="q">“{escape(r.get("quoted_text", "")[:900])}”</p></div>')
    miss = "".join(f"<li><b>{escape(m['category'])}</b> — {escape(m['why'])}</li>" for m in missing) \
        or "<li>None — every protection on the checklist was detected.</li>"
    html = f"""
    <h1>Clausify — contract risk report</h1>
    <p class="meta">{escape(meta['source'])} · {escape(meta['mode'])} · generated {datetime.now():%Y-%m-%d %H:%M}</p>
    <p class="warn">Research prototype — not legal advice. Risk levels are per clause category, scored for the
    party with less bargaining power. Models can miss clauses and
    quote the wrong text. The chance shown is calibrated on held-out CUAD contracts; clauses marked
    Possible — check have less than a 50% chance of being present.</p>
    <p><b>Summary:</b> {counts['High']} High · {counts['Medium']} Medium · {counts['Low']} Low
    {f"· {counts['Unrecognized']} Unrecognized" if counts['Unrecognized'] else ""}</p>
    <h2>Clauses, highest risk first</h2>{''.join(cards) or '<p>No clauses detected.</p>'}
    <h2>Missing-protection checklist</h2><p class="meta">Not detected in this contract (the models can miss
    clauses, so check before relying on this):</p><ul>{miss}</ul>
    """
    css = """
    body { font-family: sans-serif; font-size: 10pt; color: #222; }
    h1 { font-size: 18pt; margin-bottom: 2pt; } h2 { font-size: 13pt; margin-top: 14pt; }
    .meta { color: #666; font-size: 8.5pt; } .warn { color: #8a5a00; font-size: 8.5pt; }
    .card { padding-left: 6pt; margin: 8pt 0; } .h { margin: 0; } .why { color: #555; margin: 1pt 0; }
    .pe { margin: 2pt 0; } .sc { color: #666; font-size: 8pt; margin: 1pt 0; }
    .q { font-style: italic; color: #333; font-size: 9pt; margin: 2pt 0; }
    """
    out = io.BytesIO()
    writer = fitz.DocumentWriter(out)
    story = fitz.Story(html=html, user_css=css)
    page, where = fitz.paper_rect("a4"), fitz.paper_rect("a4") + (40, 40, -40, -40)
    more = True
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return out.getvalue()
